#!/usr/bin/env python3
from pathlib import Path
import tempfile
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

from format_aware_workbook import normalize_xlsx


with tempfile.TemporaryDirectory() as directory:
    path = Path(directory) / "fixture.xlsx"
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Providers"
    first["A1"] = "NPI"
    first["B1"] = "Name"
    first["A2"] = "1234567890"
    first["B2"] = "Example Clinic"
    first["B2"].number_format = "@"
    first["C1"] = "Formula check"
    first["C2"] = "=1+1"
    second = workbook.create_sheet("Phones")
    second["C4"] = "(555) 010-1234"
    second["C4"].hyperlink = "https://example.test/clinic"
    second.merge_cells("D5:E5")
    second["D5"] = "Merged heading"
    workbook.save(path)
    # openpyxl deliberately does not calculate formula results. Patch the
    # standards-compliant cached <v> exactly as a producer such as Excel or
    # LibreOffice would, so data_only=True exercises the available-cache path.
    with zipfile.ZipFile(path, "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    sheet_name = "xl/worksheets/sheet1.xml"
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    root = ET.fromstring(members[sheet_name])
    formula_cell = root.find(f".//{{{namespace}}}c[@r='C2']")
    cached = formula_cell.find(f"{{{namespace}}}v")
    if cached is None:
        cached = ET.SubElement(formula_cell, f"{{{namespace}}}v")
    cached.text = "2"
    members[sheet_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, body in members.items():
            archive.writestr(name, body)
    normalized = normalize_xlsx(path)
    assert normalized["normalizationFormat"] == "workbook_rows_v1"
    assert normalized["semanticFiltering"] is False
    assert normalized["completeNonemptyCellCoverage"] is True
    assert normalized["sheetCount"] == 2
    assert 'A2="1234567890"' in normalized["evidenceText"]
    assert 'B2="Example Clinic"' in normalized["evidenceText"]
    assert 'C4="(555) 010-1234"' in normalized["evidenceText"]
    assert 'C2={"formula":"=1+1","cachedDisplayedValue":2}' in normalized["evidenceText"]
    assert normalized["formulaCells"] == 1
    assert normalized["cachedFormulaValues"] == 1
    assert sum(sheet["nonemptyCells"] for sheet in normalized["sheets"]) == 8
    assert normalized["mergedCellCount"] == 1
    assert normalized["hyperlinkCount"] == 1
    assert any(row["coordinate"] == "B2" for row in normalized["nonGeneralNumberFormats"])
    assert "MERGED_CELL_RECORD_BOUNDARY" in normalized["manualReviewSignals"]
    assert "HYPERLINK_OR_COMMENT_NEEDED" in normalized["manualReviewSignals"]

print("format-aware workbook tests passed")
