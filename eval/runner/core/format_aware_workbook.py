#!/usr/bin/env python3
"""Losslessly project every nonempty XLSX cell to auditable row text."""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
from itertools import zip_longest
import json
from pathlib import Path
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter


def json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def normalize_xlsx(path: Path) -> dict:
    raw = path.read_bytes()
    with zipfile.ZipFile(path) as archive:
        xml_members = {name: archive.read(name) for name in archive.namelist()
                       if name.startswith("xl/worksheets/") and name.endswith(".xml")}
        merged_cell_count = sum(body.count(b"<mergeCell ") for body in xml_members.values())
        hyperlink_count = sum(body.count(b"<hyperlink ") for body in xml_members.values())
        hidden_record_signal_count = sum(body.count(b'hidden="1"') for body in xml_members.values())
        comment_part_count = sum(name.startswith("xl/comments") for name in archive.namelist())
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    displayed_workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = [f"# Complete workbook text", f"Workbook SHA-256: {hashlib.sha256(raw).hexdigest()}"]
    sheets = []
    non_general_formats = []
    for worksheet in workbook.worksheets:
        displayed_worksheet = displayed_workbook[worksheet.title]
        lines.extend(["", f"## Sheet: {worksheet.title}", f"Declared dimensions: {worksheet.calculate_dimension()}"])
        nonempty_rows = 0
        nonempty_cells = 0
        formula_cells = 0
        cached_formula_values = 0
        first_row = None
        last_row = None
        formula_rows = worksheet.iter_rows()
        displayed_rows = displayed_worksheet.iter_rows()
        for row_number, (row, displayed_row) in enumerate(
                zip_longest(formula_rows, displayed_rows, fillvalue=()), start=1):
            cells = []
            for cell, displayed_cell in zip_longest(row, displayed_row, fillvalue=None):
                if cell is None:
                    raise ValueError(f"DATA_ONLY_WORKBOOK_HAS_EXTRA_CELL:{worksheet.title}:{row_number}")
                if (displayed_cell is not None and hasattr(displayed_cell, "coordinate")
                        and displayed_cell.coordinate != cell.coordinate):
                    raise ValueError(
                        f"FORMULA_CACHE_COORDINATE_MISMATCH:{worksheet.title}:{cell.coordinate}:"
                        f"{displayed_cell.coordinate}")
                if cell.value is None:
                    continue
                coordinate = f"{get_column_letter(cell.column)}{cell.row}"
                value = cell.value
                if cell.number_format and cell.number_format != "General":
                    non_general_formats.append({"sheet": worksheet.title, "coordinate": coordinate,
                                                "numberFormat": cell.number_format})
                if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
                    cached = displayed_cell.value if displayed_cell is not None else None
                    formula_cells += 1
                    cached_formula_values += cached is not None
                    value = {"formula": value, "cachedDisplayedValue": json_value(cached)}
                cells.append(f"{coordinate}={json.dumps(json_value(value), ensure_ascii=False, separators=(',', ':'))}")
            if not cells:
                continue
            nonempty_rows += 1
            nonempty_cells += len(cells)
            first_row = row_number if first_row is None else first_row
            last_row = row_number
            lines.append(f"ROW {row_number}: " + " | ".join(cells))
        sheets.append({
            "name": worksheet.title,
            "declaredDimensions": worksheet.calculate_dimension(),
            "nonemptyRows": nonempty_rows,
            "nonemptyCells": nonempty_cells,
            "formulaCells": formula_cells,
            "cachedFormulaValues": cached_formula_values,
            "firstNonemptyRow": first_row,
            "lastNonemptyRow": last_row,
        })
    workbook.close()
    displayed_workbook.close()
    text = "\n".join(lines) + "\n"
    return {
        "normalizationFormat": "workbook_rows_v1",
        "extractionMode": "openpyxl_complete_nonempty_cells",
        "bodySha256": hashlib.sha256(raw).hexdigest(),
        "bodyBytes": len(raw),
        "evidenceText": text,
        "evidenceTextChars": len(text),
        "evidenceTextBytes": len(text.encode("utf-8")),
        "evidenceTextSha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "sheetCount": len(sheets),
        "sheets": sheets,
        "semanticFiltering": False,
        "completeNonemptyCellCoverage": True,
        "formulaPolicy": "formula_and_cached_displayed_value_when_available",
        "cachedValueLookupPolicy": "formula_and_data_only_read_only_rows_iterated_in_coordinate_lockstep",
        "formulaCells": sum(sheet["formulaCells"] for sheet in sheets),
        "cachedFormulaValues": sum(sheet["cachedFormulaValues"] for sheet in sheets),
        "displayFidelity": "stored_values_formulas_cached_values_and_number_format_metadata_not_excel_rendered_strings",
        "nonGeneralNumberFormats": non_general_formats,
        "mergedCellCount": merged_cell_count,
        "hyperlinkCount": hyperlink_count,
        "commentPartCount": comment_part_count,
        "hiddenRecordSignalCount": hidden_record_signal_count,
        "manualReviewSignals": [
            *(["MERGED_CELL_RECORD_BOUNDARY"] if merged_cell_count else []),
            *(["HYPERLINK_OR_COMMENT_NEEDED"] if hyperlink_count or comment_part_count else []),
            *(["AMBIGUOUS_WORKBOOK_RECORD_BOUNDARY"] if hidden_record_signal_count else []),
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    if args.workbook.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SystemExit("UNSUPPORTED_WORKBOOK_FORMAT: only .xlsx/.xlsm are supported by this pinned parser")
    print(json.dumps(normalize_xlsx(args.workbook), ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
